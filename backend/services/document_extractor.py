from __future__ import annotations
from pathlib import Path
from typing import Any


def extract_text(raw: bytes, filename: str) -> tuple[str, dict[str, Any]]:
    name = (filename or "document").lower()
    suffix = Path(name).suffix
    if suffix in {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".xml", ".html", ".htm"}:
        return raw.decode("utf-8", errors="ignore"), {"parser": "plain-text", "format": suffix.lstrip(".")}
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            import io
            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
            return text, {"parser": "pypdf", "format": "pdf", "pages": len(reader.pages)}
        except Exception as exc:
            raise ValueError(f"PDF extraction failed: {exc}") from exc
    if suffix == ".docx":
        try:
            from docx import Document
            import io
            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs)
            return text, {"parser": "python-docx", "format": "docx"}
        except Exception as exc:
            raise ValueError(f"DOCX extraction failed: {exc}") from exc
    if suffix == ".pptx":
        try:
            from pptx import Presentation
            import io
            prs = Presentation(io.BytesIO(raw))
            parts=[]
            for i, slide in enumerate(prs.slides, 1):
                parts.append(f"[Slide {i}]")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
            return "\n".join(parts), {"parser": "python-pptx", "format": "pptx", "slides": len(prs.slides)}
        except Exception as exc:
            raise ValueError(f"PPTX extraction failed: {exc}") from exc
    raise ValueError(f"Unsupported document format '{suffix or 'unknown'}'. Supported: TXT, MD, CSV, JSON, YAML, XML, HTML, PDF, DOCX, PPTX.")
